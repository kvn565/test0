"""Moteur d'importation des factures historiques (fichier Excel importation.xlsx).

Les factures importées sont créées exactement comme les factures envoyées à l'OBR
(statut_obr = 'ENVOYE', identifiant OBR, numéro, date d'envoi) sans aucun appel
réseau. Les numéros déjà présents en base sont ignorés.

Principe de calcul
------------------
On n'écrit jamais les totaux « bruts » du fichier : on alimente le moteur de
calcul du projet (Produit + LigneFacture + Facture) pour que la facture créée
soit lisible exactement comme les autres factures.

Pour chaque ligne du fichier :
    montant HT  = P.T HTTC          (la TC du fichier est déjà incluse)
    taux TVA    = TVA / P.T HTTC × 100
    taux PFL    = PF / P.T HTTC × 100
On reproduit ainsi à l'identique les montants historiques du fichier :
    P.T HTTC + TVA + PF = PT TTC
Le prix unitaire enregistré (produit et ligne) est P.T HTTC / Qtés. La TC n'est
pas appliquée une seconde fois (elle est déjà contenue dans le HT).
Chaque ligne conserve son taux TVA réel du fichier (0 %, 10 % ou 18 %), même si
le produit possède un autre taux principal.
"""

import logging
import re
import unicodedata
from collections import Counter, OrderedDict
from datetime import datetime, date
from datetime import time as dt_time
from decimal import Decimal, ROUND_DOWN, ROUND_HALF_UP

from django.db import transaction
from django.utils import timezone as dj_timezone

from categories.models import Categorie
from clients.models import Client
from facturer.models import Facture, LigneFacture
from produits.models import Produit
from societe.models import Societe
from taux.models import TauxTVA, TauxTC, TauxPFL

from .models import Utilisateur

logger = logging.getLogger(__name__)

ZERO = Decimal('0')
CENT = Decimal('100')
THREE = Decimal('0.001')
FOUR = Decimal('0.0001')

PAIEMENT_MAP = {
    'banque': 'BANQUE',
    'cash': 'CAISSE',
    'caisse': 'CAISSE',
    'espece': 'CAISSE',
    'credit': 'CREDIT',
    'autres': 'AUTRES',
    'autre': 'AUTRES',
}

STATUTS_VALIDES = ('valide', 'valides')

COLONNES_OBLIGATOIRES = (
    'Numero',
    'Type de factures',
    'Produit',
    'Qtés',
    'P.T HTTC',
    'TVA',
    'PT TTC',
    'Date',
    'Client',
)

COLONNES_OPTIONNELLES = (
    'Type de paiement',
    'Status',
    'Unité',
    'P.U HTTC',
    'TC',
    'PF',
    'OTT',
    'TSCE',
    'NIF du Client',
    'Devise',
    'Auteur',
)


class ImportationError(Exception):
    pass


def _norm(value):
    if value is None:
        return ''
    text = str(value).strip().lower()
    text = unicodedata.normalize('NFKD', text)
    text = ''.join(c for c in text if not unicodedata.combining(c))
    return re.sub(r'\s+', ' ', text)


def _clean(value):
    if value is None:
        return ''
    return re.sub(r'\s+', ' ', str(value)).strip()


def _to_decimal(value):
    if value is None:
        return ZERO
    if isinstance(value, Decimal):
        return value
    if isinstance(value, bool):
        return ZERO
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        return Decimal(str(value))
    text = str(value).strip()
    if not text:
        return ZERO
    text = text.replace('\u00a0', ' ').replace('\n', ' ').strip()
    text = text.replace(' ', '')
    if ',' in text and '.' in text:
        text = text.replace('.', '').replace(',', '.')
    else:
        text = text.replace(',', '.')
    text = re.sub(r'[^0-9.\-]', '', text)
    if not text or text in ('-', '.', '-.'):
        return ZERO
    try:
        return Decimal(text)
    except Exception:
        return ZERO


def _parse_datetime(value):
    if isinstance(value, datetime):
        return value.date(), value.time().replace(microsecond=0)
    if isinstance(value, date):
        return value, dt_time(0, 0)
    text = _clean(value)
    for fmt in (
        '%Y-%m-%d %H:%M:%S',
        '%Y-%m-%d %H:%M',
        '%d/%m/%Y %H:%M:%S',
        '%d/%m/%Y %H:%M',
        '%d-%m-%Y %H:%M:%S',
        '%d-%m-%Y %H:%M',
        '%Y-%m-%d',
        '%d/%m/%Y',
        '%d-%m-%Y',
    ):
        try:
            parsed = datetime.strptime(text, fmt)
            return parsed.date(), parsed.time().replace(microsecond=0)
        except ValueError:
            continue
    return None, None


def _index(header, *candidates, exact=False):
    for candidate in candidates:
        for position, name in enumerate(header):
            if (name == candidate) if exact else (candidate in name):
                return position
    return None


def _cell(row, index):
    if index is None or index >= len(row):
        return None
    return row[index]


def _taux(montant, base):
    """Pourcentage (3 décimales) d'un montant par rapport à une base."""
    if not base:
        return ZERO
    valeur = (montant / base * CENT).quantize(THREE, rounding=ROUND_HALF_UP)
    return valeur if valeur > 0 else ZERO


def _dominer(valeurs):
    """Valeur la plus fréquente (départage : la plus grande)."""
    valeurs = [v for v in valeurs if v is not None]
    if not valeurs:
        return ZERO
    compteur = Counter(valeurs)
    return max(compteur.items(), key=lambda item: (item[1], item[0]))[0]


def _read_rows(path):
    try:
        import openpyxl
    except ImportError as exc:
        raise ImportationError(
            "La librairie openpyxl est requise pour lire les fichiers Excel."
        ) from exc

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook['Table 1'] if 'Table 1' in workbook.sheetnames else workbook.worksheets[0]
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    return rows


def parse_fichier(path):
    rows = _read_rows(path)

    header_pos = None
    header = []
    for position, row in enumerate(rows):
        line = [_norm(cell) for cell in row]
        if 'numero' in line and any('client' in cell for cell in line):
            header_pos = position
            header = line
            break

    if header_pos is None:
        raise ImportationError(
            "En-tête du fichier introuvable (colonnes « Numero » et « Client » absentes)."
        )

    i_numero = _index(header, 'numero', exact=True)
    i_type = _index(header, 'type de factures', 'type de facture')
    i_paiement = _index(header, 'type de paiement')
    i_status = _index(header, 'status', 'statut')
    i_produit = _index(header, 'produit', exact=True)
    i_unite = _index(header, 'unite', exact=True)
    i_qte = _index(header, 'qtes', 'qte', 'quantite')
    i_pu_ht = _index(header, 'p.u httc', exact=True)
    i_pt_ht = _index(header, 'p.t httc', exact=True)
    i_tc = _index(header, 'tc', exact=True)
    i_tva = _index(header, 'tva', exact=True)
    i_pt_ttc = _index(header, 'pt ttc', exact=True)
    i_pf = _index(header, 'pf', exact=True)
    i_ott = _index(header, 'ott', exact=True)
    i_tsce = _index(header, 'tsce', exact=True)
    i_date = _index(header, 'date', exact=True)
    i_auteur = _index(header, 'auteur')
    i_client = _index(header, 'client', exact=True)
    i_nif = _index(header, 'nif')
    i_devise = _index(header, 'devise')

    manquantes = [
        nom for nom, index in (
            ('Numero', i_numero),
            ('Type de factures', i_type),
            ('Produit', i_produit),
            ('Qtés', i_qte),
            ('P.T HTTC', i_pt_ht),
            ('TVA', i_tva),
            ('PT TTC', i_pt_ttc),
            ('Date', i_date),
            ('Client', i_client),
        ) if index is None
    ]
    if manquantes:
        raise ImportationError(
            "Colonnes obligatoires manquantes : " + ", ".join(manquantes) + "."
        )

    invoices = OrderedDict()
    skipped_status = []
    errors = []
    warnings = []
    lines_total = 0
    a_ott = False
    a_tsce = False

    for row in rows[header_pos + 1:]:
        if not any(cell is not None and str(cell).strip() != '' for cell in row):
            continue

        numero = _clean(_cell(row, i_numero))
        if not numero or numero.upper().startswith('TOTAL'):
            continue

        status = _norm(_cell(row, i_status))
        if status and status not in STATUTS_VALIDES:
            if numero not in skipped_status:
                skipped_status.append(numero)
            continue

        designation = _clean(_cell(row, i_produit))
        if not designation:
            errors.append({'numero': numero, 'message': "Ligne sans désignation de produit."})
            continue

        quantite = _to_decimal(_cell(row, i_qte))
        if quantite == 0:
            quantite = Decimal('1')

        pu_ht = _to_decimal(_cell(row, i_pu_ht))
        pt_ht = _to_decimal(_cell(row, i_pt_ht))
        tc = _to_decimal(_cell(row, i_tc))
        tva = _to_decimal(_cell(row, i_tva))
        pt_ttc = _to_decimal(_cell(row, i_pt_ttc))
        pf = _to_decimal(_cell(row, i_pf))
        ott = _to_decimal(_cell(row, i_ott))
        tsce = _to_decimal(_cell(row, i_tsce))

        if ott:
            a_ott = True
        if tsce:
            a_tsce = True

        # Option B : le montant HT enregistré est P.T HTTC (TC incluse) afin de
        # reproduire exactement les montants du fichier (TVA portée sur P.T HTTC).
        base = pt_ht.quantize(THREE)
        if base <= 0:
            base = ((pu_ht * quantite) + tc).quantize(THREE)
        if base <= 0 and pt_ttc:
            base = pt_ttc.quantize(THREE)
        if base <= 0:
            errors.append({'numero': numero, 'message': "Ligne sans montant HT exploitable."})
            continue

        pu_unitaire = (base / quantite).quantize(THREE, rounding=ROUND_DOWN)

        # La TVA du fichier est calculée sur P.T HTTC ; la TC est déjà incorporée
        # dans le HT, elle n'est donc pas appliquée une seconde fois.
        taux_tc = ZERO
        taux_tva = _taux(tva, base)
        taux_pfl = _taux(pf, base)

        if numero not in invoices:
            jour, heure = _parse_datetime(_cell(row, i_date))
            type_facture = (_clean(_cell(row, i_type)) or 'FN').upper()[:2]
            if type_facture not in ('FN', 'FA'):
                type_facture = 'FN'
            invoices[numero] = {
                'numero': numero,
                'date': jour,
                'heure': heure,
                'type_facture': type_facture,
                'mode_paiement': PAIEMENT_MAP.get(_norm(_cell(row, i_paiement)), 'AUTRES'),
                'client_nom': _clean(_cell(row, i_client)) or 'CLIENT DIVERS',
                'client_nif': _clean(_cell(row, i_nif)),
                'auteur': _clean(_cell(row, i_auteur)),
                'devise': (_clean(_cell(row, i_devise)) or 'BIF').upper()[:5],
                'lignes': [],
                'total_ht_excel': ZERO,
                'total_tc_excel': ZERO,
                'total_tva_excel': ZERO,
                'total_pfl_excel': ZERO,
                'total_ttc_excel': ZERO,
            }

        facture = invoices[numero]
        facture['lignes'].append({
            'designation': designation,
            'quantite': quantite,
            'unite': _clean(_cell(row, i_unite)) or 'Unité',
            'pu_ht': pu_unitaire,
            'taux_tc': taux_tc,
            'taux_tva': taux_tva,
            'taux_pfl': taux_pfl,
            'base': base,
            'pt_ht': pt_ht,
            'pt_ttc': pt_ttc,
        })
        facture['total_ht_excel'] += base
        facture['total_tc_excel'] += tc
        facture['total_tva_excel'] += tva
        facture['total_pfl_excel'] += pf
        facture['total_ttc_excel'] += pt_ttc if pt_ttc else (base + tc + tva + pf)
        lines_total += 1

    # Taux unique par produit (valeur dominante) + prix unitaire dominant
    produits = OrderedDict()
    for facture in invoices.values():
        for ligne in facture['lignes']:
            key = ligne['designation'].upper()
            entry = produits.setdefault(key, {
                'designation': ligne['designation'],
                'unite': ligne['unite'],
                'pu_ht': [],
                'taux_tc': [],
                'taux_tva': [],
                'taux_pfl': [],
            })
            entry['pu_ht'].append(ligne['pu_ht'])
            entry['taux_tc'].append(ligne['taux_tc'])
            entry['taux_tva'].append(ligne['taux_tva'])
            entry['taux_pfl'].append(ligne['taux_pfl'])

    produits_rates = {}
    for key, entry in produits.items():
        produits_rates[key] = {
            'designation': entry['designation'],
            'unite': entry['unite'],
            'pu_ht': _dominer(entry['pu_ht']),
            'taux_tc': _dominer(entry['taux_tc']),
            'taux_tva': _dominer(entry['taux_tva']),
            'taux_pfl': _dominer(entry['taux_pfl']),
        }

    for facture in invoices.values():
        for champ in ('total_ht_excel', 'total_tc_excel', 'total_tva_excel', 'total_pfl_excel', 'total_ttc_excel'):
            facture[champ] = facture[champ].quantize(THREE, rounding=ROUND_DOWN)

    assujetti = {
        'tva': any(l['taux_tva'] > 0 for f in invoices.values() for l in f['lignes']),
        'tc': any(l['taux_tc'] > 0 for f in invoices.values() for l in f['lignes']),
        'pfl': any(l['taux_pfl'] > 0 for f in invoices.values() for l in f['lignes']),
    }

    if a_ott:
        warnings.append("Des lignes contiennent une taxe OTT non nulle ; l'application ne gère pas encore l'OTT (ignorée).")
    if a_tsce:
        warnings.append("Des lignes contiennent une taxe TSCE non nulle ; l'application ne gère pas encore la TSCE (ignorée).")

    return {
        'invoices': list(invoices.values()),
        'produits_rates': produits_rates,
        'assujetti': assujetti,
        'lines': lines_total,
        'skipped_status': skipped_status,
        'errors': errors,
        'warnings': warnings,
    }


def _build_identifier(societe, moment, numero):
    system_id = (getattr(societe, 'obr_system_id', '') or '').strip() or 'IMPORT'
    nif = (getattr(societe, 'nif', '') or '').strip() or 'NIF'
    return f"{nif}/{system_id}/{moment.strftime('%Y%m%d%H%M%S')}/{numero}"


def _set_mode(model, instance, mode):
    if instance.obr_mode_envoye != mode:
        model.objects.filter(pk=instance.pk).update(obr_mode_envoye=mode)
        instance.obr_mode_envoye = mode


def _activer_assujettissement(societe, assujetti):
    """Active sur la société les taxes détectées dans le fichier."""
    changements = {}
    for champ, cle in (('assujeti_tva', 'tva'), ('assujeti_tc', 'tc'), ('assujeti_pfl', 'pfl')):
        if assujetti.get(cle) and not getattr(societe, champ, False):
            changements[champ] = True
    if changements:
        Societe.objects.filter(pk=societe.pk).update(**changements)
        for champ, valeur in changements.items():
            setattr(societe, champ, valeur)
    return changements


def _resoudre_categorie(societe, mode):
    categorie = Categorie.objects.filter(societe=societe, obr_mode_envoye=mode).first()
    if categorie:
        return categorie
    categorie = Categorie.objects.filter(societe=societe).first()
    if categorie:
        return categorie
    categorie = Categorie(societe=societe, nom='Importation')
    categorie.save()
    _set_mode(Categorie, categorie, mode)
    return categorie


def _get_or_create_client(societe, mode, nom, nif, user):
    nom = _clean(nom) or 'CLIENT DIVERS'
    nif = _clean(nif) or None

    qs = Client.objects.filter(societe=societe)
    client = None
    if nif:
        client = qs.filter(nif=nif).first()
    if client is None:
        client = qs.filter(nom__iexact=nom).first()
    if client is not None:
        return client, False

    client = Client(societe=societe, nom=nom, nif=nif, cree_par=user)
    client.save()
    _set_mode(Client, client, mode)
    return client, True


def _get_or_create_taux(model, societe, mode, valeur, label):
    valeur = Decimal(valeur).quantize(THREE)
    taux = model.objects.filter(societe=societe, valeur=valeur).first()
    if taux is not None:
        return taux, False

    nom = f"{label} {valeur}%"
    if model.objects.filter(societe=societe, nom=nom).exists():
        compteur = 1
        while model.objects.filter(societe=societe, nom=f"{nom} ({compteur})").exists():
            compteur += 1
        nom = f"{nom} ({compteur})"

    taux = model(societe=societe, nom=nom, valeur=valeur, est_defaut=False)
    taux.save()
    _set_mode(model, taux, mode)
    return taux, True


def _get_or_create_produit(societe, mode, designation, rate, devise, categorie):
    designation = _clean(designation)
    produit = Produit.objects.filter(societe=societe, designation__iexact=designation).first()
    if produit is not None:
        return produit, False

    produit = Produit(
        societe=societe,
        categorie=categorie,
        designation=designation,
        unite=rate.get('unite') or 'Unité',
        prix_vente=rate['pu_ht'] or ZERO,
        devise=devise or 'BIF',
        taux_tva=rate['taux_tva_obj'],
        taux_tc=rate['taux_tc_obj'],
        taux_pfl=rate['taux_pfl_obj'],
        origine='LOCAL',
        statut='ACTIF',
    )
    produit.save()
    _set_mode(Produit, produit, mode)
    return produit, True


def _creer_facture(societe, mode, facture_data, client, user):
    moment = dj_timezone.make_aware(
        datetime.combine(facture_data['date'], facture_data['heure']),
        dj_timezone.get_current_timezone(),
    )
    identifier = _build_identifier(societe, moment, facture_data['numero'])

    cree_par = user
    auteur = facture_data.get('auteur')
    if auteur:
        utilisateur = Utilisateur.objects.filter(societe=societe, username__iexact=auteur).first()
        if utilisateur is not None:
            cree_par = utilisateur

    facture = Facture(
        societe=societe,
        numero=facture_data['numero'],
        invoice_identifier=identifier,
        date_facture=facture_data['date'],
        heure_facture=facture_data['heure'],
        type_facture=facture_data['type_facture'],
        client=client,
        devise=facture_data['devise'],
        mode_paiement=facture_data['mode_paiement'],
        applique_tva=True,
        statut_obr='ENVOYE',
        message_obr="Importé depuis un fichier Excel (facture historique) — statut OBR : envoyé.",
        date_envoi_obr=moment,
        obr_registered_date=moment,
        obr_registered_number=facture_data['numero'],
        cree_par=cree_par,
    )
    facture.save()
    _set_mode(Facture, facture, mode)
    return facture


def _creer_ligne(facture, produit, ligne, taux_tva):
    """Crée la ligne via le moteur du projet (prix TVAC → HT/TVA).

    Le taux passé est le taux réel de la ligne historique (0, 10 ou 18 % selon
    le fichier). Comme le modèle force normalement le taux du produit, on
    réapplique ensuite le taux de la ligne pour les factures importées.
    """
    taux_valeur = taux_tva.valeur if taux_tva else ZERO
    quantite = ligne['quantite'] or Decimal('1')
    base = (ligne.get('base') or (ligne['pu_ht'] * quantite)).quantize(THREE)
    pu_ht = (base / quantite).quantize(FOUR, rounding=ROUND_DOWN)

    if taux_valeur == 0:
        prix_tvac = (base / quantite).quantize(THREE, rounding=ROUND_HALF_UP)
    else:
        prix_tvac = (
            (base * (Decimal('1') + taux_valeur / CENT)) / quantite
        ).quantize(THREE, rounding=ROUND_HALF_UP)

    ligne_obj = LigneFacture(
        facture=facture,
        produit=produit,
        designation=_clean(ligne['designation']),
        quantite=ligne['quantite'],
        prix_vente_tvac=prix_tvac,
        taux_tva=taux_tva,
    )
    ligne_obj.save()

    LigneFacture.objects.filter(pk=ligne_obj.pk).update(
        taux_tva=taux_tva,
        prix_vente_tvac=prix_tvac,
        prix_unitaire_ht=pu_ht,
    )
    for attr in ('_montant_ht', '_montant_tva', '_montant_ttc'):
        if hasattr(ligne_obj, attr):
            delattr(ligne_obj, attr)
    ligne_obj.taux_tva = taux_tva
    ligne_obj.prix_vente_tvac = prix_tvac
    ligne_obj.prix_unitaire_ht = pu_ht
    return ligne_obj


def _index_produits(societe):
    return {_clean(p.designation).upper() for p in Produit.objects.filter(societe=societe)}


def _index_clients(societe):
    noms = {_clean(c.nom).upper() for c in Client.objects.filter(societe=societe)}
    nifs = {_clean(c.nif) for c in Client.objects.filter(societe=societe) if c.nif}
    return noms, nifs


def analyser_fichier(path, societe, mode):
    parsed = parse_fichier(path)

    existing_numbers = set(
        Facture.objects.filter(societe=societe).values_list('numero', flat=True)
    )

    duplicates = []
    to_import = []
    for facture in parsed['invoices']:
        if facture['numero'] in existing_numbers:
            duplicates.append(facture['numero'])
        else:
            to_import.append(facture)

    clients_map = OrderedDict()
    for facture in to_import:
        nom = _clean(facture['client_nom']) or 'CLIENT DIVERS'
        clients_map.setdefault(nom.upper(), (nom, _clean(facture['client_nif'])))

    existing_client_names, existing_client_nifs = _index_clients(societe)
    existing_produits = _index_produits(societe)

    nouveaux_clients = []
    for key, (nom, nif) in clients_map.items():
        if key in existing_client_names:
            continue
        if nif and nif in existing_client_nifs:
            continue
        nouveaux_clients.append(nom)

    produits_utilises = OrderedDict()
    for facture in to_import:
        for ligne in facture['lignes']:
            key = ligne['designation'].upper()
            produits_utilises.setdefault(key, ligne['designation'])
    nouveaux_produits = [
        designation for key, designation in produits_utilises.items() if key not in existing_produits
    ]

    total_ttc_excel = sum((f['total_ttc_excel'] for f in to_import), ZERO)

    apercu = [
        {
            'numero': f['numero'],
            'date': f['date'],
            'client': f['client_nom'],
            'nb_lignes': len(f['lignes']),
            'total_ttc_excel': f['total_ttc_excel'],
        }
        for f in to_import[:100]
    ]

    return {
        'societe_id': societe.pk,
        'societe_nom': societe.nom,
        'mode_label': 'PRODUCTION' if mode else 'TEST',
        'invoices_found': len(parsed['invoices']),
        'invoices_to_import': len(to_import),
        'lines_found': parsed['lines'],
        'total_ttc_excel': total_ttc_excel.quantize(THREE, rounding=ROUND_DOWN),
        'assujetti': parsed['assujetti'],
        'assujettissement': [cle for cle, actif in parsed['assujetti'].items() if actif],
        'duplicates': duplicates,
        'skipped_status': parsed['skipped_status'],
        'errors': parsed['errors'],
        'warnings': parsed['warnings'],
        'nouveaux_clients': nouveaux_clients,
        'nouveaux_produits': nouveaux_produits,
        'apercu': apercu,
    }


def importer_fichier(path, societe, mode, user):
    parsed = parse_fichier(path)

    changements_societe = _activer_assujettissement(societe, parsed['assujetti'])

    report = {
        'societe_id': societe.pk,
        'societe_nom': societe.nom,
        'mode_label': 'PRODUCTION' if mode else 'TEST',
        'invoices_found': len(parsed['invoices']),
        'lines_found': parsed['lines'],
        'created': 0,
        'skipped_existing': 0,
        'skipped_status': len(parsed['skipped_status']),
        'clients_created': 0,
        'produits_created': 0,
        'taux_created': 0,
        'assujettissement': list(changements_societe.keys()),
        'errors': [],
        'warnings': list(parsed['warnings']),
    }

    produits_rates = parsed['produits_rates']
    taux_cache = {}

    def _taux_obj(model, valeur, label):
        valeur = Decimal(valeur).quantize(THREE)
        cle = (model.__name__, valeur)
        if cle in taux_cache:
            return taux_cache[cle]
        obj, created = _get_or_create_taux(model, societe, mode, valeur, label)
        if created:
            report['taux_created'] += 1
        taux_cache[cle] = obj
        return obj

    categorie = None
    for facture_data in parsed['invoices']:
        numero = facture_data['numero']
        if Facture.objects.filter(societe=societe, numero=numero).exists():
            report['skipped_existing'] += 1
            continue

        if facture_data['date'] is None or facture_data['heure'] is None:
            report['errors'].append({'numero': numero, 'message': "Date illisible."})
            continue

        try:
            with transaction.atomic():
                client, client_created = _get_or_create_client(
                    societe, mode, facture_data['client_nom'], facture_data['client_nif'], user
                )
                if client_created:
                    report['clients_created'] += 1

                if categorie is None:
                    categorie = _resoudre_categorie(societe, mode)

                facture = _creer_facture(societe, mode, facture_data, client, user)

                for ligne in facture_data['lignes']:
                    key = ligne['designation'].upper()
                    rate = produits_rates.get(key, {
                        'designation': ligne['designation'],
                        'unite': ligne['unite'],
                        'pu_ht': ligne['pu_ht'],
                        'taux_tc': ligne['taux_tc'],
                        'taux_tva': ligne['taux_tva'],
                        'taux_pfl': ligne['taux_pfl'],
                    })
                    rate = dict(rate)
                    rate['taux_tva_obj'] = _taux_obj(TauxTVA, rate['taux_tva'], 'TVA')
                    rate['taux_tc_obj'] = (
                        _taux_obj(TauxTC, rate['taux_tc'], 'TC') if rate['taux_tc'] > 0 else None
                    )
                    rate['taux_pfl_obj'] = (
                        _taux_obj(TauxPFL, rate['taux_pfl'], 'PFL') if rate['taux_pfl'] > 0 else None
                    )

                    produit, produit_created = _get_or_create_produit(
                        societe, mode, ligne['designation'], rate,
                        facture_data['devise'], categorie,
                    )
                    if produit_created:
                        report['produits_created'] += 1

                    taux_ligne = _taux_obj(TauxTVA, ligne['taux_tva'], 'TVA')
                    _creer_ligne(facture, produit, ligne, taux_ligne)

                facture.recalculer_totaux()

            report['created'] += 1

        except Exception as exc:
            logger.exception("[Importation] Échec facture %s", numero)
            report['errors'].append({'numero': numero, 'message': str(exc)[:300]})

    return report
